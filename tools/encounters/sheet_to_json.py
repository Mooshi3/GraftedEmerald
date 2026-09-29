#!/usr/bin/env python3
"""Build src/data/wild_encounters.json (Emerald maps) from the encounter spreadsheet.

The spreadsheet (tools/encounters/Grafted_Encounter_Tables.xlsx) is the source of truth
for which Pokémon appear where. This script:
  - resolves every name in the sheet to a species and checks it (unknown names, legendaries,
    evolved forms far below their evolution level, too many surf/fishing entries),
  - writes species into the matching maps' tables, keeping each slot's existing levels for
    land (water rows carry their own level range in the "Levels" column),
  - makes sure every Emerald map keeps at least the encounter tables it had in vanilla
    (fishing, rock smash...), copying vanilla data where the sheet has nothing,
  - generates src/data/graft_postgame_legendaries.h from the "Postgame Legendaries" tab.

Usage:  python3 tools/encounters/sheet_to_json.py [--check]
        --check only validates and reports; nothing is written.
Needs:  pip install openpyxl
"""
import collections
import glob
import json
import os
import re
import subprocess
import sys

import openpyxl

ROOT = os.path.normpath(os.path.join(os.path.dirname(__file__), '..', '..'))
SHEET = os.path.join(ROOT, 'tools', 'encounters', 'Grafted_Encounter_Tables.xlsx')
JSON_OUT = os.path.join(ROOT, 'src', 'data', 'wild_encounters.json')
LEGENDS_OUT = os.path.join(ROOT, 'src', 'data', 'graft_postgame_legendaries.h')
VANILLA_REV = '21c2420299'  # upstream expansion commit Grafted branched from

LAND_ZONES = {'Grass', 'Cave', 'Desert'}
TIMES = ('Morning', 'Day', 'Night')
ROD_SLOTS = {'Old Rod': (0, 2), 'Good Rod': (2, 5), 'Super Rod': (5, 10)}

# Sheet area name -> Emerald maps it fills. Land rows fill every listed map;
# water/fishing rows fill the listed maps that have water.
AREA_MAPS = {
    'MT. PYRE EXTERIOR': ['MT_PYRE_EXTERIOR', 'MT_PYRE_SUMMIT'],
    'MT. PYRE INTERIOR': ['MT_PYRE_1F', 'MT_PYRE_2F', 'MT_PYRE_3F', 'MT_PYRE_4F', 'MT_PYRE_5F', 'MT_PYRE_6F'],
    'AQUA HIDEOUT': ['AQUA_HIDEOUT_1F'],
    'MAGMA HIDEOUT': ['MAGMA_HIDEOUT_1F', 'MAGMA_HIDEOUT_2F_1R', 'MAGMA_HIDEOUT_2F_2R', 'MAGMA_HIDEOUT_2F_3R',
                      'MAGMA_HIDEOUT_3F_1R', 'MAGMA_HIDEOUT_3F_2R', 'MAGMA_HIDEOUT_3F_3R', 'MAGMA_HIDEOUT_4F'],
    'NEW MAUVILLE': ['NEW_MAUVILLE_ENTRANCE', 'NEW_MAUVILLE_INSIDE'],
    'SAFARI ZONE': ['SAFARI_ZONE_SOUTH', 'SAFARI_ZONE_SOUTHWEST', 'SAFARI_ZONE_NORTH',
                    'SAFARI_ZONE_NORTHWEST', 'SAFARI_ZONE_SOUTHEAST', 'SAFARI_ZONE_NORTHEAST'],
    'SAFARI ZONE SW': ['SAFARI_ZONE_SOUTHWEST'],
    'SAFARI ZONE NW': ['SAFARI_ZONE_NORTHWEST'],
    'SAFARI ZONE SE': ['SAFARI_ZONE_SOUTHEAST'],
    'SHOAL CAVE': ['SHOAL_CAVE_LOW_TIDE_ENTRANCE_ROOM', 'SHOAL_CAVE_LOW_TIDE_STAIRS_ROOM', 'SHOAL_CAVE_LOW_TIDE_LOWER_ROOM',
                   'SHOAL_CAVE_LOW_TIDE_INNER_ROOM', 'SHOAL_CAVE_LOW_TIDE_ICE_ROOM'],
    'SEAFLOOR CAVERN': ['SEAFLOOR_CAVERN_ENTRANCE'] + [f'SEAFLOOR_CAVERN_ROOM{i}' for i in range(1, 9)],
    'CAVE OF ORIGIN': ['CAVE_OF_ORIGIN_ENTRANCE', 'CAVE_OF_ORIGIN_1F', 'CAVE_OF_ORIGIN_UNUSED_RUBY_SAPPHIRE_MAP1',
                       'CAVE_OF_ORIGIN_UNUSED_RUBY_SAPPHIRE_MAP2', 'CAVE_OF_ORIGIN_UNUSED_RUBY_SAPPHIRE_MAP3'],
    'METEOR FALLS': ['METEOR_FALLS_1F_1R', 'METEOR_FALLS_1F_2R', 'METEOR_FALLS_B1F_1R', 'METEOR_FALLS_B1F_2R'],
    'METEOR FALLS 1F': ['METEOR_FALLS_1F_1R', 'METEOR_FALLS_1F_2R'],
    'METEOR FALLS B1F': ['METEOR_FALLS_B1F_1R', 'METEOR_FALLS_B1F_2R'],
    'ABANDONED SHIP': ['ABANDONED_SHIP_ROOMS_B1F', 'ABANDONED_SHIP_HIDDEN_FLOOR_CORRIDORS'],
    'UNDERWATER 124': ['UNDERWATER_ROUTE124'],
    'UNDERWATER 126': ['UNDERWATER_ROUTE126'],
}


def area_maps(area):
    if area in AREA_MAPS:
        return AREA_MAPS[area]
    name = re.sub(r'[^A-Z0-9]+', '_', area.upper().replace('.', '')).strip('_')
    return [re.sub(r'^ROUTE_(\d+)$', r'ROUTE\1', name)]


# ---------------------------------------------------------------- species data
def load_species():
    info, bydex, evo_level = {}, collections.defaultdict(list), {}
    for path in sorted(glob.glob(os.path.join(ROOT, 'src/data/pokemon/species_info/gen_*_families.h'))):
        text = open(path, encoding='utf-8').read()
        for m in re.finditer(r'\[(SPECIES_[A-Z0-9_]+)\]\s*=\s*\{(.*?)\n    \},', text, re.S):
            body = m.group(2)
            dex = re.search(r'\.natDexNum\s*=\s*(NATIONAL_DEX_\w+)', body)
            flags = set(re.findall(r'\.(is\w+)\s*=\s*TRUE', body))
            info[m.group(1)] = dict(dex=dex.group(1) if dex else None, flags=flags)
            if dex:
                bydex[dex.group(1)].append(m.group(1))
        for m in re.finditer(r'\{EVO_LEVEL[A-Z_]*,\s*(\d+),\s*(SPECIES_[A-Z0-9_]+)', text):
            evo_level[m.group(2)] = max(evo_level.get(m.group(2), 0), int(m.group(1)))
    return info, bydex, evo_level


SPECIES, BYDEX, EVO_LEVEL = load_species()
LEGENDARY = {'isLegendary', 'isRestrictedLegendary', 'isSubLegendary', 'isMythical', 'isUltraBeast', 'isParadox'}
# Names whose default form has its own species id.
FORM_NAMES = {'CASTFORM': 'CASTFORM_NORMAL', 'SINISTEA': 'SINISTEA_PHONY', 'POLTCHAGEIST': 'POLTCHAGEIST_COUNTERFEIT',
              'MIMIKYU': 'MIMIKYU_DISGUISED', 'DEERLING': 'DEERLING_SPRING', 'EISCUE': 'EISCUE_ICE',
              'WISHIWASHI': 'WISHIWASHI_SOLO', 'BASCULIN': 'BASCULIN_RED_STRIPED', 'TATSUGIRI': 'TATSUGIRI_CURLY',
              'FLOETTE': 'FLOETTE_RED', 'MORPEKO': 'MORPEKO_FULL_BELLY', 'PUMPKABOO': 'PUMPKABOO_AVERAGE',
              'SQUAWKABILLY': 'SQUAWKABILLY_GREEN', 'ORICORIO': 'ORICORIO_BAILE', 'SHELLOS': 'SHELLOS_WEST',
              'BURMY': 'BURMY_PLANT', 'INDEEDEE': 'INDEEDEE_M', 'GIRATINA': 'GIRATINA_ALTERED',
              'TORNADUS': 'TORNADUS_INCARNATE', 'THUNDURUS': 'THUNDURUS_INCARNATE', 'LANDORUS': 'LANDORUS_INCARNATE',
              'ENAMORUS': 'ENAMORUS_INCARNATE', 'ZYGARDE': 'ZYGARDE_50', 'XERNEAS': 'XERNEAS_NEUTRAL',
              'ZACIAN': 'ZACIAN_HERO', 'ZAMAZENTA': 'ZAMAZENTA_HERO', 'TERAPAGOS': 'TERAPAGOS_NORMAL',
              'FLABEBE': 'FLABEBE_RED', 'SCATTERBUG': 'SCATTERBUG', 'FURFROU': 'FURFROU_NATURAL',
              'MINIOR': 'MINIOR_METEOR_RED', 'UNOWN': 'UNOWN'}


def resolve(name):
    n = name.strip().replace('♀', ' F').replace('♂', ' M').replace('é', 'e').replace('É', 'E')
    key = re.sub(r'[^A-Z0-9]+', '_', n.upper().replace('.', '').replace("'", '').replace(':', '')).strip('_')
    key = {'TAILOW': 'TAILLOW'}.get(key, key)
    key = FORM_NAMES.get(key, key)
    if 'SPECIES_' + key in SPECIES:
        return 'SPECIES_' + key
    forms = BYDEX.get('NATIONAL_DEX_' + key)
    if forms:
        return forms[0]
    return 'SPECIES_' + key  # species defined through macros (Unown, Flabébé...) still have this id


def species_exists(sid):
    if sid in SPECIES:
        return True
    return subprocess.run(['grep', '-rqw', sid, os.path.join(ROOT, 'include/constants/species.h')]).returncode == 0


# ---------------------------------------------------------------- map tiles
def _map_tile_index():
    import struct
    beh_c = open(os.path.join(ROOT, 'src/metatile_behavior.c'), encoding='utf-8').read()
    encounter_behaviors = set(re.findall(r'\[(MB_\w+)\]\s*=\s*TILE_FLAG_[A-Z_| ]*HAS_ENCOUNTERS', beh_c))
    water = {'MB_POND_WATER', 'MB_INTERIOR_DEEP_WATER', 'MB_DEEP_WATER', 'MB_SOOTOPOLIS_DEEP_WATER', 'MB_OCEAN_WATER',
             'MB_NO_SURFACING', 'MB_SEAWEED', 'MB_SEAWEED_NO_SURFACING', 'MB_WATERFALL', 'MB_EASTWARD_CURRENT',
             'MB_WESTWARD_CURRENT', 'MB_NORTHWARD_CURRENT', 'MB_SOUTHWARD_CURRENT', 'MB_UNUSED_SOOTOPOLIS_DEEP_WATER',
             'MB_UNUSED_SOOTOPOLIS_DEEP_WATER_2', 'MB_SHOAL_CAVE_ENTRANCE'}
    names, value = {}, 0
    for n, v in re.findall(r'^\s*(MB_[A-Z0-9_]+)(?:\s*=\s*(0x[0-9A-Fa-f]+|\d+))?\s*,',
                           open(os.path.join(ROOT, 'include/constants/metatile_behaviors.h'), encoding='utf-8').read(), re.M):
        value = int(v, 0) if v else value
        names[value] = n
        value += 1
    layouts = {l['id']: l for l in json.load(open(os.path.join(ROOT, 'data/layouts/layouts.json')))['layouts'] if 'id' in l}
    tileset_dirs = {}
    for kind in ('primary', 'secondary'):
        for d in os.listdir(os.path.join(ROOT, 'data/tilesets', kind)):
            tileset_dirs[d.replace('_', '').lower()] = os.path.join(ROOT, 'data/tilesets', kind, d)
    def attrs(symbol):
        path = tileset_dirs.get(re.sub(r'^gTileset_', '', symbol).replace('_', '').lower())
        if not path or not os.path.exists(path + '/metatile_attributes.bin'):
            return []
        raw = open(path + '/metatile_attributes.bin', 'rb').read()
        return [struct.unpack_from('<H', raw, i)[0] for i in range(0, len(raw), 2)]
    index = {}
    for mapdir in glob.glob(os.path.join(ROOT, 'data/maps/*/map.json')):
        m = json.load(open(mapdir, encoding='utf-8'))
        layout = layouts.get(m.get('layout'))
        if not layout or not layout.get('blockdata_filepath'):
            continue
        blocks = open(os.path.join(ROOT, layout['blockdata_filepath']), 'rb').read()
        primary, secondary = attrs(layout['primary_tileset']), attrs(layout['secondary_tileset'])
        land = water_tiles = False
        for i in range(0, len(blocks), 2):
            mid = struct.unpack_from('<H', blocks, i)[0] & 0x3FF
            table, idx = (primary, mid) if mid < 512 else (secondary, mid - 512)
            behavior = names.get(table[idx] & 0xFF) if idx < len(table) else None
            if behavior in encounter_behaviors:
                if behavior in water:
                    water_tiles = True
                else:
                    land = True
        index[m['id'][4:]] = (land, water_tiles)
    return index


MAP_TILES = _map_tile_index()
# Tiles placed at runtime, so the layout scan can't see them: Mirage Island's grass on Route 130.
MAP_TILES['ROUTE130'] = (True, True)


# ---------------------------------------------------------------- sheet
def read_sheet(problems, warnings):
    wb = openpyxl.load_workbook(SHEET, data_only=True)
    ws = wb['Encounter Tables']
    header = [str(c.value or '').strip().lower() for c in ws[1]]
    col_levels = header.index('levels') if 'levels' in header else None
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0] or not str(r[0]).strip():
            continue
        area, zone, time = str(r[0]).strip(), str(r[1] or '').strip(), str(r[2] or '').strip()
        mons = [str(x).strip() if x is not None else '' for x in r[3:15]]
        lv = None
        if col_levels is not None and r[col_levels]:
            m = re.search(r'(\d+)\s*-\s*(\d+)', str(r[col_levels]))
            lv = (int(m.group(1)), int(m.group(2))) if m else None
        rows.append(dict(area=area, zone=zone, time=time, mons=mons, levels=lv))
    legends = []
    if 'Postgame Legendaries' in wb.sheetnames:
        for r in wb['Postgame Legendaries'].iter_rows(min_row=2, values_only=True):
            if r[0] and r[3]:
                legends.append(dict(area=str(r[0]).strip(), zone=str(r[1]).strip(), time=str(r[2]).strip(),
                                    name=str(r[3]).strip()))
    return rows, legends


def check_names(where, names, problems, warnings, levels=None, allow_legendary=False):
    out = []
    for n in names:
        if not n:
            out.append(None)
            continue
        sid = resolve(n)
        if not species_exists(sid):
            problems.append(f'{where}: unknown species "{n}" ({sid})')
        elif sid in SPECIES and SPECIES[sid]['flags'] & LEGENDARY and not allow_legendary:
            problems.append(f'{where}: {n} is legendary (legendaries belong on the Postgame Legendaries tab)')
        need = EVO_LEVEL.get(sid, 0)
        if levels and need and need > levels[1] + 10:
            warnings.append(f'{where}: {n} normally evolves at Lv {need}, area tops out at Lv {levels[1]}')
        out.append(sid)
    return out


# ---------------------------------------------------------------- json helpers
def emerald(entries):
    return [e for e in entries if 'map' in e and not re.search('FireRed|LeafGreen', e['base_label'])]


def entry_time(e):
    m = re.search(r'_(Morning|Day|Evening|Night)$', e['base_label'])
    return m.group(1) if m else None


def find_entry(entries, mapname, time):
    """Day and untimed are the same table (untimed labels count as the fallback time, Day)."""
    for e in entries:
        if e.get('map') == 'MAP_' + mapname and not re.search('FireRed|LeafGreen', e['base_label']):
            t = entry_time(e)
            if t == time or (time == 'Day' and t is None):
                return e
    return None


def map_entries(entries, mapname):
    return [e for e in emerald(entries) if e['map'] == 'MAP_' + mapname]


def set_species(field, species):
    for slot, sid in zip(field['mons'], species):
        slot['species'] = sid


def main():
    check_only = '--check' in sys.argv
    problems, warnings, notes = [], [], []
    rows, legends = read_sheet(problems, warnings)
    data = json.load(open(JSON_OUT, encoding='utf-8'))
    group = data['wild_encounter_groups'][0]
    entries = group['encounters']
    vanilla = json.loads(subprocess.run(['git', '-C', ROOT, 'show', f'{VANILLA_REV}:src/data/wild_encounters.json'],
                                        capture_output=True, text=True, check=True).stdout)
    vanilla_entries = emerald(vanilla['wild_encounter_groups'][0]['encounters'])

    # 1. Every map keeps at least the tables it had in vanilla (in its Day/untimed entry).
    for ve in vanilla_entries:
        mapname = ve['map'][4:]
        target = find_entry(entries, mapname, 'Day')
        if target is None:
            target = json.loads(json.dumps(ve))
            entries.append(target)
            notes.append(f'{mapname}: whole map was missing; restored from vanilla')
            continue
        for field in ('land_mons', 'water_mons', 'rock_smash_mons', 'fishing_mons'):
            if field in ve and field not in target and not (field == 'land_mons' and map_entries(entries, mapname)):
                target[field] = json.loads(json.dumps(ve[field]))
                notes.append(f'{mapname}: restored missing {field} from vanilla')
    # A land table on a map with no land encounter tiles (grass, cave floor, sand...) never triggers.
    for e in emerald(entries):
        if 'land_mons' in e and not MAP_TILES.get(e['map'][4:], (True, True))[0]:
            del e['land_mons']
            notes.append(f"{e['map'][4:]}: removed land table (the map has no tiles that trigger land encounters)")
    group['encounters'] = entries = [e for e in entries if any(k.endswith('_mons') for k in e)]

    # 2. Apply sheet rows.
    covered = set()
    fishing = collections.defaultdict(dict)
    for row in rows:
        where = f"{row['area']} / {row['time'] or row['zone']}"
        maps = area_maps(row['area'])
        if row['zone'] in LAND_ZONES:
            species = check_names(where, row['mons'], problems, warnings)
            if None in species:
                problems.append(f'{where}: land tables need all 12 slots filled ({species.count(None)} empty)')
                continue
            times = TIMES if row['time'] in ('', 'All') else (row['time'],)
            for mapname in maps:
                existing = map_entries(entries, mapname)
                if not MAP_TILES.get(mapname, (True, True))[0]:
                    warnings.append(f'{where}: {mapname} has no tiles that trigger land encounters; this row does nothing there')
                    continue
                if not existing or not any('land_mons' in e for e in existing):
                    problems.append(f'{where}: map {mapname} has no land table to fill')
                    continue
                for t in times:
                    e = find_entry(entries, mapname, t)
                    if e is None:
                        if row['time'] in ('', 'All'):
                            continue  # untimed row: fill only the tables the map has
                        base = find_entry(entries, mapname, 'Day') or existing[0]
                        e = {'map': base['map'], 'base_label': re.sub(r'_(Morning|Day|Evening|Night)$', '', base['base_label']) + '_' + t,
                             'land_mons': json.loads(json.dumps(base['land_mons']))}
                        entries.insert(entries.index(base) + 1, e)
                    if 'land_mons' not in e:
                        problems.append(f'{where}: {mapname} has no grass/cave table')
                        continue
                    set_species(e['land_mons'], species)
                    covered.add((mapname, 'land_mons'))
        elif row['zone'] == 'Water':
            if row['time'].startswith('Fishing') or row['area'] in ('Old Rod', 'Good Rod', 'Super Rod'):
                continue  # retired global rod rows
            if not row['levels']:
                problems.append(f'{where}: water rows need a level range in the Levels column')
                continue
            names = [n for n in row['mons'] if n]
            rod = next((r for r in ROD_SLOTS if row['time'].startswith(r)), None)
            if rod:
                lo, hi = ROD_SLOTS[rod]
                if len(names) != hi - lo:
                    problems.append(f'{where}: {rod} needs exactly {hi - lo} Pokémon, found {len(names)}')
                    continue
                for mapname in maps:
                    fishing[mapname][rod] = (check_names(where, names, problems, warnings, row['levels']), row['levels'])
            else:
                if len(names) > 5:
                    warnings.append(f'{where}: surfing has 5 slots; dropping {", ".join(names[5:])}')
                    names = names[:5]
                if len(names) < 5:
                    problems.append(f'{where}: surfing needs 5 Pokémon, found {len(names)}')
                    continue
                species = check_names(where, names, problems, warnings, row['levels'])
                for mapname in maps:
                    e = find_entry(entries, mapname, 'Day')
                    if e is None or 'water_mons' not in e:
                        if len(maps) == 1:
                            problems.append(f'{where}: {mapname} has no surfing table')
                        continue
                    set_species(e['water_mons'], species)
                    for slot in e['water_mons']['mons']:
                        slot['min_level'], slot['max_level'] = row['levels']
                    covered.add((mapname, 'water_mons'))
    for mapname, rods in fishing.items():
        e = find_entry(entries, mapname, 'Day')
        if e is None or 'fishing_mons' not in e:
            if len(area_maps(mapname)) == 1 and not any(mapname in v for v in AREA_MAPS.values()):
                problems.append(f'{mapname}: has no fishing table')
            continue
        if set(rods) != set(ROD_SLOTS):
            problems.append(f'{mapname}: fishing needs Old, Good and Super Rod rows (has {", ".join(sorted(rods))})')
            continue
        for rod, (species, (lo, hi)) in rods.items():
            a, b = ROD_SLOTS[rod]
            for slot, sid in zip(e['fishing_mons']['mons'][a:b], species):
                slot['species'], slot['min_level'], slot['max_level'] = sid, lo, hi
        covered.add((mapname, 'fishing_mons'))

    # 3. Report what is still untouched by the sheet.
    for e in emerald(entries):
        for field in ('land_mons', 'water_mons', 'fishing_mons'):
            if field in e and (e['map'][4:], field) not in covered:
                notes.append(f"{e['base_label']}: {field} not in the sheet (kept as is)")
        for field in ('land_mons', 'water_mons', 'fishing_mons', 'rock_smash_mons'):
            if field in e and any(m['species'] == 'SPECIES_UNOWN' for m in e[field]['mons']) \
                    and not e['map'].startswith('MAP_CAVE_OF_ORIGIN'):
                warnings.append(f"{e['base_label']}: {field} still has Unown placeholders")

    # 4. Postgame legendaries.
    legend_rows = []
    for L in legends:
        sid = check_names(f"Postgame / {L['area']}", [L['name']], problems, warnings, allow_legendary=True)[0]
        if sid in SPECIES and not SPECIES[sid]['flags'] & LEGENDARY:
            problems.append(f"Postgame / {L['area']}: {L['name']} is not a legendary")
        kind = 'WILD_AREA_WATER' if L['zone'] == 'Water' else 'WILD_AREA_LAND'
        time = 'TIMES_OF_DAY_COUNT' if L['time'] in ('All', 'Surfing', '') else 'TIME_' + L['time'].upper()
        for mapname in area_maps(L['area']):
            if kind == 'WILD_AREA_WATER' and not ((f := find_entry(entries, mapname, 'Day')) and 'water_mons' in f):
                continue
            legend_rows.append(f'    {{ MAP_{mapname}, {time}, {kind}, {sid} }},')

    for w in warnings:
        print('warning:', w)
    for n in notes:
        print('note:', n)
    if problems:
        for p in problems:
            print('ERROR:', p)
        print(f'{len(problems)} error(s); nothing written.')
        sys.exit(1)
    if check_only:
        print('Sheet OK (check only; nothing written).')
        return
    with open(JSON_OUT, 'w', encoding='utf-8') as f:
        f.write(json.dumps(data, indent=4, ensure_ascii=False))
    with open(LEGENDS_OUT, 'w', encoding='utf-8') as f:
        f.write('// Generated by tools/encounters/sheet_to_json.py from the "Postgame Legendaries" tab. Do not edit.\n')
        f.write('// { map, time of day (TIMES_OF_DAY_COUNT = any), area, species }\n')
        f.write('static const struct PostgameLegendary sPostgameLegendaries[] =\n{\n')
        f.write('\n'.join(legend_rows) + '\n};\n')
    print(f'Wrote {os.path.relpath(JSON_OUT, ROOT)} and {os.path.relpath(LEGENDS_OUT, ROOT)} '
          f'({len(legend_rows)} legendary slots).')


if __name__ == '__main__':
    main()
